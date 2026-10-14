"""Regression tests for exact university logos, template spillover and footer preservation."""
import io
import unittest
from zipfile import ZipFile
from docx import Document
from standalone.verified_metadata_v5 import _fill, _remove_trailing_liz_placeholders, _protected_artifacts


class UniversityBrandingTests(unittest.TestCase):
    def test_logo_and_footer_hashes_preserved(self):
        def package(logo, header=b"HEADER", footer=b"FOOTER"):
            b=io.BytesIO()
            with ZipFile(b, "w") as z:
                z.writestr("word/document.xml", "<w:document/>")
                z.writestr("word/media/UA-LOGO.png", logo)
                z.writestr("word/header1.xml", header)
                z.writestr("word/footer1.xml", footer)
            return b.getvalue()
        x=package(b"original")
        self.assertEqual(_protected_artifacts(x),_protected_artifacts(package(b"original")))
        self.assertNotEqual(_protected_artifacts(x),_protected_artifacts(package(b"altered")))

    def test_equivalent_header_xml_serialization_is_safe(self):
        def pkg(header):
            b=io.BytesIO()
            with ZipFile(b,"w") as z:
                z.writestr("word/media/image1.png",b"REAL-BRAND-BYTES")
                z.writestr("word/header1.xml",header)
            return b.getvalue()
        left=pkg(b'<w:hdr xmlns:w="urn:word"><w:p /></w:hdr>')
        right=pkg(b'<w:hdr xmlns:w="urn:word"><w:p/></w:hdr>')
        self.assertEqual(_protected_artifacts(left),_protected_artifacts(right))

    def test_missing_university_logo_fails_closed(self):
        b=io.BytesIO()
        with ZipFile(b,"w") as z:z.writestr("word/document.xml", "<w:document/>")
        with self.assertRaises(ValueError):
            _protected_artifacts(b.getvalue())

    def test_fill_deletes_old_instruction_paragraphs(self):
        doc=Document()
        t=doc.add_table(rows=1,cols=2)
        cell=t.cell(0,1)
        cell.paragraphs[0].add_run("Ejemplo anterior")
        cell.add_paragraph("Indicación de Liz")
        cell.add_paragraph("Marcador antiguo")
        _fill(cell,"Dato real verificado")
        self.assertEqual(len(cell.paragraphs),1)
        self.assertEqual(cell.text,"Dato real verificado")

    def test_trailing_blank_page_removed_without_tampering_with_table(self):
        d=Document()
        d.add_paragraph("HOJA DE METADATOS")
        t=d.add_table(rows=1,cols=2)
        t.cell(0,1).text="Dato verificado"
        d.add_paragraph("")
        d.add_paragraph("")
        self.assertEqual(_remove_trailing_liz_placeholders(d),2)
        self.assertEqual(len(d.paragraphs),1)
        self.assertEqual(len(d.tables),1)

    def test_never_remove_real_post_table_content(self):
        d=Document()
        d.add_table(rows=1,cols=2)
        d.add_paragraph("Nota legal válida")
        with self.assertRaises(ValueError):
            _remove_trailing_liz_placeholders(d)

if __name__=="__main__":
    unittest.main()
