"""Pruebas unitarias para generación de muestra y lectura Google Drive."""
from io import BytesIO
from unittest import TestCase
from unittest.mock import patch

from docx import Document
from standalone.demo_v6 import demo_metadata
from standalone.drive_adapter_v6 import (
    connection_status, list_folder_tree, DriveConnectionError,
    FOLDER_MIME,
)


class DemoWordTests(TestCase):
    def test_real_word_is_valid_and_explicitly_provisional(self):
        payload = demo_metadata(2)
        self.assertEqual(payload[:2], b"PK")
        doc = Document(BytesIO(payload))
        text = "\n".join(
            [p.text for p in doc.paragraphs] +
            [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
        )
        self.assertIn("NO VÁLIDO", text)
        self.assertIn("AUTOR DE PRUEBA 2", text)
        self.assertIn("FIRMAS", text.upper())

    def test_only_1_to_3_demo_authors(self):
        for n in [0, 4, 10]:
            with self.assertRaises(ValueError):
                demo_metadata(n)


class DriveAdapterTests(TestCase):
    def test_disabled_without_credentials(self):
        with patch.dict("os.environ", {}, clear=True):
            self.assertEqual(connection_status()["status"], "CREDENTIALS_NOT_CONFIGURED")
            with self.assertRaises(DriveConnectionError):
                list_folder_tree("1abcdefghi_123456")

    def test_incorrect_folder_id_blocked_early(self):
        with self.assertRaises(DriveConnectionError):
            list_folder_tree("../../etc/passwd")

    def test_readonly_inventory(self):
        class Response:
            status_code=200
            def __init__(self, payload):
                self.payload=payload
            def raise_for_status(self):
                pass
            def json(self):
                return self.payload
        class FakeSession:
            def __init__(self):
                self.calls=[]
            def get(self,url,params,timeout):
                self.calls.append((url, params, timeout))
                if "/files/" in url:
                    return Response({"id":"1abcdefghij123456", "mimeType":FOLDER_MIME})
                return Response({"files":[
                    {"id":"1file_abcdef123456","name":"TI.docx",
                     "mimeType":"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
                    {"id":"1file_gggggg123456","name":"Declaración.pdf","mimeType":"application/pdf"}
                ]})
        fake=FakeSession()
        result=list_folder_tree("1abcdefghij123456",session=fake)
        self.assertEqual(result["file_count"], 2)
        self.assertFalse(result["content_verified"])
        self.assertFalse(result["identity_linkage_verified"])
        self.assertTrue(result["read_only"])
        self.assertEqual(len(fake.calls), 2)
