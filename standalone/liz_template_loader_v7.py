"""Plantillas originales de Liz: recuperación verificada desde Drive.

Nunca se reconstruye un documento oficial, ni se utiliza una plantilla legacy
como sustitución silenciosa. Los archivos no se publican en GitHub.
"""
from __future__ import annotations
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from .drive_adapter_v6 import _session, DriveConnectionError

TEMPLATES = {
    1: {"file_id": "1p6-lmC8TPSdTbxAnaWTMimzlBTZrM1Vs",
        "sha256": "68cce83a4d16072a44b47a96b59ed0abd544441f45ecc6b8227406950f82e46a",
        "rows": 29},
    2: {"file_id": "1FC1P3AWieqXy6wWDyEBl-JDTYpIlMZE0",
        "sha256": "388531bf40c5340253201a72b19d419785353b8e1081f5c70144c0e40dbb1579",
        "rows": 34},
    3: {"file_id": "1aMnypc6CbFExT5g9q_OTO8CW69damZHS",
        "sha256": "a8c5122abcf73637da02c89c4ed0581131a1f1f6acd9619f667cae084c8b6a41",
        "rows": 39},
}
SOURCE = "Liz · 1 | FORMATOS CON LOGO NUEVO / 2 | POSGRADO / 1 | Maestría"


def validate_official_template(payload, author_count):
    from docx import Document
    template = TEMPLATES.get(author_count)
    if not template:
        raise ValueError("Solo se han validado formatos Liz para 1, 2 y 3 autores de maestría")
    if not isinstance(payload, bytes) or not 10000 <= len(payload) <= 2_000_000:
        raise ValueError("Bytes de plantilla oficial ausentes o fuera de límites")
    if sha256(payload).hexdigest() != template["sha256"]:
        raise ValueError("La plantilla de Liz cambió: requiere nueva validación institucional antes de usarla")
    try:
        with ZipFile(BytesIO(payload)) as archive:
            if archive.testzip() is not None or "word/document.xml" not in archive.namelist():
                raise ValueError("Archivo DOCX oficial no íntegro")
        document = Document(BytesIO(payload))
    except (BadZipFile, KeyError) as exc:
        raise ValueError("Plantilla oficial no es un DOCX válido") from exc
    if len(document.tables) != 1 or len(document.tables[0].rows) != template["rows"]:
        raise ValueError("La estructura de tablas no corresponde al formato original de Liz")
    if "HOJA DE METADATOS" not in "\n".join(p.text for p in document.paragraphs).upper():
        raise ValueError("Encabezado institucional no corresponde a hoja de metadatos")
    return True


def fetch_official_metadata_template(author_count, session=None):
    meta = TEMPLATES.get(author_count)
    if not meta:
        raise ValueError("Número de autores fuera del catálogo Liz")
    session = session or _session()
    try:
        r = session.get(
            "https://www.googleapis.com/drive/v3/files/" + meta["file_id"],
            params={"alt": "media", "supportsAllDrives": "true"},
            timeout=(7, 30)
        )
        if r.status_code in (401, 403, 404):
            raise DriveConnectionError("El servidor no tiene permiso para leer la plantilla oficial de Liz")
        r.raise_for_status()
        payload = r.content
    except DriveConnectionError:
        raise
    except Exception as exc:
        raise DriveConnectionError("No se pudo leer la plantilla oficial desde Drive") from exc
    validate_official_template(payload, author_count)
    return payload
