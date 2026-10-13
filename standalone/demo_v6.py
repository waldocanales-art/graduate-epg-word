"""Documento demostrativo sin datos personales para probar la generación DOCX.

NO reproduce ni sustituye la plantilla oficial de Liz.
"""
from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

def demo_metadata(author_count=2):
    if author_count not in (1, 2, 3):
        raise ValueError("El ejemplo acepta entre 1 y 3 autores")
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.65)
    sec.bottom_margin = Inches(.65)
    sec.left_margin = Inches(.75)
    sec.right_margin = Inches(.75)
    header = sec.header.paragraphs[0]
    header.text = "UA | HAWKS     ·     ESCUELA DE POSGRADO"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        run.font.size = Pt(9)
        run.bold = True
        run.font.color.rgb = RGBColor(188, 83, 17)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("HOJA DE METADATOS — DEMOSTRACIÓN")
    run.bold = True
    run.font.size = Pt(15)
    doc.add_paragraph("EJEMPLO CON DATOS FICTICIOS. NO VÁLIDO PARA FIRMA, SUSTENTACIÓN NI REPOSITORIO.")
    doc.add_paragraph("Prueba técnica del generador HAWKS. No es la plantilla oficial institucional.")
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"

    def section(label):
        cells = table.add_row().cells
        cells[0].merge(cells[1])
        cells[0].text = label
        for p in cells[0].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(188, 83, 17)

    def row(name, value):
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = value
        for run in cells[0].paragraphs[0].runs:
            run.bold = True

    for i in range(1, author_count + 1):
        section("Autor " + str(i) + " · Datos ficticios")
        row("Nombres y apellidos", "AUTOR DE PRUEBA " + str(i))
        row("Documento de identidad", "[DNI POR VERIFICAR]")
        row("ORCID", "[ORCID POR VERIFICAR]")
    section("Asesor y jurado")
    row("Asesor", "[ASESOR POR VERIFICAR]")
    row("Jurado", "[PRESIDENTE / SECRETARIO / VOCAL POR VERIFICAR]")
    section("Investigación")
    row("Título", "EJEMPLO DE INVESTIGACIÓN PARA PRUEBA DEL SISTEMA")
    row("Programa", "[MAESTRÍA O SEGUNDA ESPECIALIDAD POR VERIFICAR]")
    row("Guía y resolución", "[POR VERIFICAR CON EXPEDIENTE]")
    row("Línea de investigación", "[POR VERIFICAR]")
    row("RSU / OCDE", "[POR VERIFICAR]")
    row("Firmas y huellas", "NO INCORPORADAS: REQUIEREN EVIDENCIA AUTÉNTICA")
    doc.add_paragraph("ESTADO: BORRADOR DEMOSTRATIVO. No implica conformidad con formatos oficiales.")
    footer = sec.footer.paragraphs[0]
    footer.text = "HAWKS Documentos LAB · Archivo generado para pruebas · Sin datos reales"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
