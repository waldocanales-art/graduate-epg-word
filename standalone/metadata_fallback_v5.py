"""Fallback DOCX layout for metadata when legacy V1 Word template is incompatible.

Always labelled preliminary, never described as the September 2026 Liz original.
"""
from io import BytesIO
from docx import Document
from docx.shared import Inches,RGBColor,Pt

def new_metadata(group_id, profile):
    people=profile["authors"]
    if not people or len(people)>3:raise ValueError("Autores incompletos")
    doc=Document();sec=doc.sections[0]
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
    sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    hd=sec.header.paragraphs[0]
    hd.text="UNIVERSIDAD AUTÓNOMA DEL PERÚ | ESCUELA DE POSGRADO"
    for r in hd.runs:r.font.bold=True;r.font.color.rgb=RGBColor(206,96,13);r.font.size=Pt(9)
    title=doc.add_heading("HOJA DE METADATOS",0);title.alignment=1
    doc.add_paragraph("Grupo EPG "+str(group_id)+" · Datos recuperados del expediente institucional",style="Subtitle")
    table=doc.add_table(rows=0,cols=2);table.style="Table Grid"
    count=0
    def section(text):
        r=table.add_row();cell=r.cells[0];cell.merge(r.cells[1]);cell.text=text
        for para in cell.paragraphs:
            for run in para.runs:run.font.bold=True;run.font.color.rgb=RGBColor(206,96,13)
    def row(label, value):
        nonlocal count
        if not value:raise ValueError("Falta dato verificado: "+label)
        r=table.add_row();r.cells[0].text=label;r.cells[1].text=str(value)
        for x in r.cells[0].paragraphs[0].runs:x.font.bold=True
        count+=1
    for i,p in enumerate(people,1):
        section("Datos del autor "+str(i))
        for label,key in [("Nombres y apellidos","name"),("Tipo de documento de identidad","document_type"),("Número de documento de identidad","document_number"),("URL de ORCID","orcid")]:row(label,p.get(key))
    section("Datos del asesor")
    for label,key in [("Nombres y apellidos","name"),("Tipo de documento de identidad","document_type"),("Número de documento de identidad","document_number"),("URL de ORCID","orcid")]:row(label,profile["adviser"].get(key))
    for label,k in [("Presidente del jurado","president"),("Secretario del jurado","secretary"),("Vocal del jurado","vocal")]:
        section(label)
        p=profile["jury"][k]
        for lb,key in [("Nombres y apellidos","name"),("Tipo de documento","document_type"),("Número de documento de identidad","document_number")]:row(lb,p.get(key))
    section("Datos de la investigación")
    for label,key in [("Título de la investigación","title"),("Línea de investigación Institucional","institutional_line"),("Línea de investigación del Programa","program_line"),("Línea de acción RSU","rsu_line"),("URL de disciplinas OCDE","ocde")]:row(label,profile["research"].get(key))
    doc.add_paragraph("BORRADOR PARA REVISIÓN INSTITUCIONAL · Diseño reconstruido; cotejar con plantilla oficial de Liz.",style="Caption")
    out=BytesIO();doc.save(out)
    return out.getvalue(),{"fields_replaced":count,"author_count":len(people),
                            "status":"GENERATED_PENDING_INSTITUTIONAL_REVIEW",
                            "template_variant":"RECONSTRUCTED_NOT_OFFICIAL_LIZ"}
