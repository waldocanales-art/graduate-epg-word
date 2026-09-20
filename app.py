from fastapi import FastAPI, HTTPException
from fastapi.responses import JSONResponse
from pydantic import BaseModel, Field
from typing import Dict, Optional, Any
from pathlib import Path
from docx import Document
from docx.shared import Cm
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
import base64, tempfile, subprocess, shutil, os, re

app = FastAPI(title='GRADUATE EPG Word Renderer', version='1.0.0')
BASE = Path(__file__).resolve().parent
TEMPLATE_DIR = Path(os.getenv('TEMPLATE_DIR', str(BASE / 'templates')))

TEMPLATES = {
    'ACTA_SUSTENTACION': '01_Plantilla_Acta_Sustentacion_WORD.docx',
    'ACTA_ORIGINALIDAD': '02_Plantilla_Acta_Aprobacion_Originalidad_WORD.docx',
    'AUTORIZACION_PUBLICACION': '03_Plantilla_Autorizacion_Publicacion_WORD.docx',
    'DECLARACION_ORIGINALIDAD': '04_Plantilla_Declaracion_Originalidad_Autor_WORD.docx',
    'DECLARACION_JURADA': '05_Plantilla_Declaracion_Jurada_Autoria_WORD.docx',
    'METADATA_1': '06_Plantilla_Hoja_Metadatos_1_Autor_WORD.docx',
    'METADATA_2': '07_Plantilla_Hoja_Metadatos_2_Autores_WORD.docx',
    'METADATA_3': '08_Plantilla_Hoja_Metadatos_3_Autores_WORD.docx',
}

IMAGE_WIDTHS_CM = {
    'IMG_AUTHOR_SIGNATURE': 4.2,
    'IMG_ADVISER_SIGNATURE': 4.2,
    'IMG_JURY_PRESIDENT_SIGNATURE': 4.0,
    'IMG_JURY_SECRETARY_SIGNATURE': 4.0,
    'IMG_JURY_VOCAL_SIGNATURE': 4.0,
    'IMG_AUTHOR1_SIGNATURE': 3.2,
    'IMG_AUTHOR2_SIGNATURE': 3.2,
    'IMG_AUTHOR3_SIGNATURE': 3.2,
    'IMG_AUTHOR1_FINGERPRINT': 2.0,
    'IMG_AUTHOR2_FINGERPRINT': 2.0,
    'IMG_AUTHOR3_FINGERPRINT': 2.0,
}


WNS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

def ooxml_replace_tokens(docx_path: Path, fields: Dict[str, Any]):
    """Replace {{FIELD}} even inside Word content controls/text boxes."""
    tmp=docx_path.with_suffix('.patched.docx')
    with ZipFile(docx_path,'r') as zin, ZipFile(tmp,'w',ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data=zin.read(item.filename)
            if item.filename.startswith('word/') and item.filename.endswith('.xml'):
                try:
                    root=etree.fromstring(data)
                    for p in root.xpath('.//w:p',namespaces=WNS):
                        texts=p.xpath('.//w:t',namespaces=WNS)
                        if not texts: continue
                        full=''.join(t.text or '' for t in texts)
                        new=full
                        for key,val in fields.items():
                            new=new.replace('{{'+str(key)+'}}','' if val is None else str(val))
                        if new!=full:
                            texts[0].text=new
                            for t in texts[1:]: t.text=''
                    data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone='yes')
                except Exception:
                    pass
            zout.writestr(item,data)
    tmp.replace(docx_path)

def unresolved_ooxml(docx_path: Path):
    found=[]
    with ZipFile(docx_path,'r') as zin:
        for name in zin.namelist():
            if name.startswith('word/') and name.endswith('.xml'):
                try:
                    txt=zin.read(name).decode('utf-8','ignore')
                    found += re.findall(r'\{\{([A-Z0-9_]+)\}\}',txt)
                except Exception:
                    pass
    return sorted(set(found))

class ImagePayload(BaseModel):
    content_type: str = 'image/png'
    base64: str

class RenderRequest(BaseModel):
    template_key: str
    group_id: Optional[str] = None
    fields: Dict[str, Any] = Field(default_factory=dict)
    images: Dict[str, ImagePayload] = Field(default_factory=dict)
    output_name: Optional[str] = None
    make_pdf: bool = True


def iter_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
                for nested in cell.tables:
                    for r in nested.rows:
                        for c in r.cells:
                            for p in c.paragraphs:
                                yield p
    for sec in doc.sections:
        for part in [sec.header, sec.footer, sec.first_page_header, sec.first_page_footer, sec.even_page_header, sec.even_page_footer]:
            for p in part.paragraphs:
                yield p
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            yield p


def collapse_replace(paragraph, replacements):
    full=''.join(r.text for r in paragraph.runs)
    if not full:
        full=paragraph.text or ''
    new=full
    for token,val in replacements.items():
        new=new.replace('{{'+token+'}}', '' if val is None else str(val))
    if new==full:
        return
    if paragraph.runs:
        for r in paragraph.runs:
            r.text=''
        paragraph.runs[0].text=new
    else:
        paragraph.add_run(new)


def replace_images(doc, images):
    for p in list(iter_paragraphs(doc)):
        full=''.join(r.text for r in p.runs) or p.text or ''
        for key,payload in images.items():
            k = key if key.startswith('IMG_') else 'IMG_'+key
            token='{{'+k+'}}'
            if token not in full:
                continue
            try:
                raw=base64.b64decode(payload.base64)
            except Exception as e:
                raise HTTPException(400, f'Imagen {key}: base64 inválido') from e
            for r in p.runs:
                r.text=''
            r=p.runs[0] if p.runs else p.add_run()
            width=Cm(IMAGE_WIDTHS_CM.get(k,3.5))
            r.add_picture(BytesIO(raw), width=width)
            return_full=(full.replace(token,'')).strip()
            if return_full:
                p.add_run(return_full)
            full=return_full


def convert_pdf(docx_path: Path, outdir: Path) -> Optional[Path]:
    soffice=shutil.which('libreoffice') or shutil.which('soffice')
    if not soffice:
        return None
    cmd=[soffice,'--headless','--convert-to','pdf','--outdir',str(outdir),str(docx_path)]
    proc=subprocess.run(cmd,stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=90)
    pdf=outdir/(docx_path.stem+'.pdf')
    return pdf if proc.returncode==0 and pdf.exists() else None

@app.get('/health')
def health():
    return {'ok':True,'templates':sorted(TEMPLATES)}

@app.post('/v1/word/render')
def render(req: RenderRequest):
    key=req.template_key.upper().strip()
    if key not in TEMPLATES:
        raise HTTPException(400, f'template_key no válido: {key}')
    template=TEMPLATE_DIR/TEMPLATES[key]
    if not template.exists():
        raise HTTPException(500, f'Plantilla no instalada: {template.name}')
    fields={str(k):v for k,v in req.fields.items()}
    out_name=req.output_name or f'{key}_{req.group_id or "DOCUMENTO"}'
    out_name=re.sub(r'[^A-Za-z0-9._-]+','_',out_name).strip('_')
    with tempfile.TemporaryDirectory() as td:
        td=Path(td)
        doc=Document(template)
        # First text, then images. IMG tokens are intentionally left intact here.
        text_fields={k:v for k,v in fields.items() if not str(k).startswith('IMG_')}
        for p in list(iter_paragraphs(doc)):
            collapse_replace(p,text_fields)
        replace_images(doc,req.images)
        docx=td/(out_name+'.docx')
        doc.save(docx)
        ooxml_replace_tokens(docx,text_fields)
        unresolved=unresolved_ooxml(docx)
        pdf=convert_pdf(docx,td) if req.make_pdf else None
        resp={
            'ok':True,
            'template_key':key,
            'group_id':req.group_id,
            'docx_filename':docx.name,
            'docx_base64':base64.b64encode(docx.read_bytes()).decode('ascii'),
            'pdf_filename':pdf.name if pdf else None,
            'pdf_base64':base64.b64encode(pdf.read_bytes()).decode('ascii') if pdf else None,
            'unresolved_tokens':sorted(set(unresolved)),
            'pdf_generated':bool(pdf),
        }
        return JSONResponse(resp)
