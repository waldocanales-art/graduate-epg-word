import tempfile
import unittest
from pathlib import Path
from docx import Document
from document_quality import audit_document


class QualityTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "sample.docx"

    def tearDown(self):
        self.tmp.cleanup()

    def test_split_run_token(self):
        d = Document()
        p = d.add_paragraph()
        p.add_run("Título {{TI")
        p.add_run("TLE}}")
        d.save(self.path)
        result = audit_document(self.path, "METADATA_1", {"TITLE": "Ejemplo"})
        self.assertFalse(result["ok"])
        self.assertEqual(result["unresolved_tokens"], ["TITLE"])

    def test_complete_document(self):
        d = Document()
        d.add_paragraph("Título final")
        d.save(self.path)
        result = audit_document(self.path, "DECLARACION_JURADA",
                                {"TITLE": "Título final", "AUTHORS_DISPLAY": "Autor"})
        self.assertTrue(result["ok"])

    def test_missing_fields(self):
        Document().save(self.path)
        result = audit_document(self.path, "DECLARACION_JURADA")
        self.assertEqual(sum(x["code"] == "MISSING_FIELD" for x in result["findings"]), 2)

    def test_invalid_zip(self):
        self.path.write_text("Not DOCX")
        self.assertFalse(audit_document(self.path, "METADATA_1", {"TITLE": "ok"})["ok"])


if __name__ == "__main__":
    unittest.main()
